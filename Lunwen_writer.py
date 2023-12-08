from docx import Document

def replace_table_cell(doc, table_index, row_index, col_index, new_text):
    try:
        table = doc.tables[table_index]
        cell = table.cell(row_index, col_index)
        cell.text = new_text
    except IndexError:
        print(f"Error: Table with index {table_index} or cell at row {row_index}, column {col_index} not found.")

if __name__ == "__main__":
    # 读取文档
    doc_path = 'path/to/your/document.docx'
    doc = Document(doc_path)

    # 要替换的表格和单元格的索引以及新文本
    table_index = 0  # 第一个表格
    row_index = 0    # 第一行
    col_index = 0    # 第一列
    new_text = 'NewValue'

    # 替换指定单元格的文本
    replace_table_cell(doc, table_index, row_index, col_index, new_text)

    # 保存修改后的文档
    modified_doc_path = 'path/to/your/modified_document.docx'
    doc.save(modified_doc_path)

    print(f"修改后的文档已保存：{modified_doc_path}")
